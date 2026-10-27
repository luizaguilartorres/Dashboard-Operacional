from flask import Flask, jsonify, request, Blueprint, session, send_file
import io
import sqlite3
from datetime import datetime
from pathlib import Path

CAMINHO = Path(__file__).parent
(CAMINHO / "db").mkdir(exist_ok=True)
CAMINHO_BANCO = CAMINHO / "db" / "dados_base_flora.db"

flora = Blueprint("flora", __name__)


@flora.before_request
def proteger_api():
    if "usuario" not in session:
        return jsonify(erro="não autenticado"), 401
    if session.get("empresa") not in ("4log", "flora"):
        return jsonify(erro="sem permissão"), 403


def data():
    ano = request.args.get("ano")  # Ex.: /api/flora/recebimento/resumo?ano=2026&mes=3 -> ano="2026"
    mes = request.args.get("mes")
    if mes:
        mes = mes.zfill(2)  # completa com zero à esquerda. Ex: 3 -> 03
    return ano, mes


def filtro_dados(coluna="data_rec"):
    ano, mes = data()
    filtro, params = [], []
    if ano:
        filtro.append(f"strftime('%Y',{coluna}) = ?")
        params.append(ano)
    if mes:
        filtro.append(f"strftime('%m',{coluna}) = ?")
        params.append(mes)
    where = "WHERE " + " AND ".join(filtro) if filtro else ""
    return where, params


def consultar(sql, params=()):
    conn = sqlite3.connect(CAMINHO_BANCO)
    conn.row_factory = sqlite3.Row
    try:
        return [dict(l) for l in conn.execute(sql, params).fetchall()]
    finally:
        conn.close()


def dias_semana(coluna):
    return f"""CASE CAST(strftime('%w', {coluna}) AS INTEGER)
            WHEN 0 THEN 'Domingo'
            WHEN 1 THEN 'Segunda'
            WHEN 2 THEN 'Terça'
            WHEN 3 THEN 'Quarta'
            WHEN 4 THEN 'Quinta'
            WHEN 5 THEN 'Sexta'
            WHEN 6 THEN 'Sábado'
            END"""


# ==========================================================
# BASE: a tabela "in_out" da Flora tem recebimento e saída na MESMA linha.
#  - Cada linha = uma carga recebida (um CT-e). "mesmo carro" na observação
#    quer dizer que o caminhão levou mais de uma carga, por isso o dashboard
#    fala em CARGAS e não em carros.
#  - Recebimento: data_recebimento. Saída (expedição): data_coleta_saida,
#    e check_saida = 1 quando a carga já saiu.
#  - A planilha vem com ~3.400 linhas vazias no final; o WHERE as descarta.
#  - data_coleta_saida é texto e tem datas digitadas fora do padrão
#    (ex.: 25/082025, 07/07//2026). Tirando "/", "-" e espaços sobra ddmmaaaa,
#    que é remontado como aaaa-mm-dd.
#  - Cliente/destino vêm com espaço no fim ("JACAREI ") -> TRIM.
#  - no_nota_fiscal pode ter várias notas separadas por vírgula
#    (às vezes com vírgula sobrando no fim) -> qtd_nf conta as notas.
# ==========================================================
BASE = """
WITH a AS (
    SELECT rowid AS id,
           date(data_recebimento)                              AS data_rec,
           UPPER(TRIM(cliente))                                AS cliente,
           rtrim(replace(COALESCE(no_nota_fiscal, ''), ' ', ''), ',') AS nf,
           CAST(no_cte AS INTEGER)                             AS cte,
           CAST(COALESCE(quantidade_paletes, 0) AS INTEGER)    AS pallets,
           UPPER(TRIM(rota_destino))                           AS destino,
           COALESCE(check_saida, 0)                            AS saiu,
           data_coleta_saida                                   AS saida_bruta,
           replace(replace(replace(TRIM(COALESCE(data_coleta_saida, '')), '/', ''), '-', ''), ' ', '') AS saida_txt,
           observacao
    FROM in_out
    WHERE data_recebimento IS NOT NULL
), b AS (
    SELECT id, data_rec, cliente, nf, cte, pallets, destino, saiu, observacao,
           CASE WHEN nf = '' THEN 0
                ELSE LENGTH(nf) - LENGTH(replace(nf, ',', '')) + 1 END AS qtd_nf,
           CASE
               WHEN saida_bruta GLOB '[0-9][0-9][0-9][0-9]-[0-9][0-9]-[0-9][0-9]*'
                    THEN date(saida_bruta)
               WHEN LENGTH(saida_txt) = 8 AND saida_txt NOT GLOB '*[^0-9]*'
                    THEN substr(saida_txt, 5, 4) || '-' || substr(saida_txt, 3, 2) || '-' || substr(saida_txt, 1, 2)
           END AS data_saida
    FROM a
)
"""

# Expedição = cargas que já saíram (e têm data de saída válida)
EXPEDIDAS = "saiu = 1 AND data_saida IS NOT NULL"

# Dias que a carga ficou armazenada (só faz sentido para carga expedida)
PERMANENCIA = "(julianday(data_saida) - julianday(data_rec))"


def rota(sql, coluna="data_rec", extra=None):
    # No SQL use {where} para o filtro de ano/mês, {col} para a coluna de data
    # usada no filtro e {dias} para o nome do dia da semana dessa coluna.
    # Não coloque "f" antes das aspas: o .format() abaixo é quem preenche.
    # "extra" é uma condição fixa que se soma ao filtro.
    where, params = filtro_dados(coluna)
    if extra:
        where = f"{where} AND ({extra})" if where else f"WHERE {extra}"
    sql = (BASE + sql).format(where=where, col=coluna, dias=dias_semana(coluna))
    return jsonify(consultar(sql, params))


def rec(sql):
    return rota(sql, "data_rec")


def exp(sql):
    return rota(sql, "data_saida", EXPEDIDAS)


# ==========================
# API RECEBIMENTO
# ==========================
@flora.route("/api/flora/recebimento/resumo")
def flora_recebimento_resumo():
    return rec("""
        SELECT COUNT(*)                                   AS cargas,
               COALESCE(SUM(pallets), 0)                  AS pallets,
               COALESCE(SUM(qtd_nf), 0)                   AS notas,
               COUNT(DISTINCT cliente)                    AS clientes,
               COUNT(DISTINCT destino)                    AS destinos,
               ROUND(COALESCE(SUM(pallets), 0) * 1.0 / NULLIF(COUNT(*), 0), 1) AS media_pallets
        FROM b
        {where}
    """)


@flora.route("/api/flora/recebimento/pallet/data")
def flora_recebimento_pallet_data():
    return rec("""
        SELECT date({col}) AS data, COALESCE(SUM(pallets), 0) AS total
        FROM b
        {where}
        GROUP BY date({col})
        ORDER BY data
    """)


@flora.route("/api/flora/recebimento/carga/data")
def flora_recebimento_carga_data():
    return rec("""
        SELECT date({col}) AS data, COUNT(*) AS total
        FROM b
        {where}
        GROUP BY date({col})
        ORDER BY data
    """)


@flora.route("/api/flora/recebimento/pallet/dia-semana")
def flora_recebimento_pallet_dia_semana():
    return rec("""
        SELECT strftime('%w', {col}) AS num_dia,
               {dias} AS dia_semana,
               COALESCE(SUM(pallets), 0) AS total
        FROM b
        {where}
        GROUP BY strftime('%w', {col})
        ORDER BY num_dia
    """)


@flora.route("/api/flora/recebimento/pallet/cliente")
def flora_recebimento_pallet_cliente():
    return rec("""
        SELECT cliente, COALESCE(SUM(pallets), 0) AS total
        FROM b
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@flora.route("/api/flora/recebimento/pallet/destino")
def flora_recebimento_pallet_destino():
    return rec("""
        SELECT destino, COALESCE(SUM(pallets), 0) AS total
        FROM b
        {where}
        GROUP BY destino
        ORDER BY total DESC
    """)


@flora.route("/api/flora/recebimento/carga/por-cliente")
def flora_recebimento_carga_cliente():
    return rec("""
        SELECT cliente, COUNT(*) AS total
        FROM b
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@flora.route("/api/flora/recebimento/nota-fiscal/cliente")
def flora_recebimento_nota_cliente():
    return rec("""
        SELECT cliente, COALESCE(SUM(qtd_nf), 0) AS total
        FROM b
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@flora.route("/api/flora/recebimento/por-mes")
def flora_recebimento_por_mes():
    return rec("""
        SELECT strftime('%Y-%m', {col}) AS mes, COUNT(*) AS cargas,
               COALESCE(SUM(pallets), 0) AS total
        FROM b
        {where}
        GROUP BY strftime('%Y-%m', {col})
        ORDER BY mes
    """)


@flora.route("/api/flora/recebimento/lista")
def flora_recebimento_lista():
    return rec("""
        SELECT data_rec AS data, cliente, nf AS nota_fiscal, cte, pallets, destino,
               data_saida AS saida,
               CASE WHEN saiu = 1 THEN 'Expedido' ELSE 'Em estoque' END AS situacao,
               CASE WHEN saiu = 1 AND data_saida IS NOT NULL
                    THEN CAST(julianday(data_saida) - julianday(data_rec) AS INTEGER) END AS dias,
               observacao
        FROM b
        {where}
        ORDER BY data_rec DESC, cliente
    """)


# ==========================
# API EXPEDIÇÃO (cargas que já saíram, filtradas pela data de saída)
# ==========================
@flora.route("/api/flora/expedicao/resumo")
def flora_expedicao_resumo():
    return exp(f"""
        SELECT COUNT(*)                                   AS cargas,
               COALESCE(SUM(pallets), 0)                  AS pallets,
               COALESCE(SUM(qtd_nf), 0)                   AS notas,
               COUNT(DISTINCT cliente)                    AS clientes,
               COUNT(DISTINCT destino)                    AS destinos,
               ROUND(AVG({PERMANENCIA}), 1)               AS permanencia_media,
               CAST(MAX({PERMANENCIA}) AS INTEGER)        AS permanencia_max
        FROM b
        {{where}}
    """)


@flora.route("/api/flora/expedicao/pallet/data")
def flora_expedicao_pallet_data():
    return exp("""
        SELECT date({col}) AS data, COALESCE(SUM(pallets), 0) AS total
        FROM b
        {where}
        GROUP BY date({col})
        ORDER BY data
    """)


@flora.route("/api/flora/expedicao/carga/data")
def flora_expedicao_carga_data():
    return exp("""
        SELECT date({col}) AS data, COUNT(*) AS total
        FROM b
        {where}
        GROUP BY date({col})
        ORDER BY data
    """)


@flora.route("/api/flora/expedicao/pallet/dia-semana")
def flora_expedicao_pallet_dia_semana():
    return exp("""
        SELECT strftime('%w', {col}) AS num_dia,
               {dias} AS dia_semana,
               COALESCE(SUM(pallets), 0) AS total
        FROM b
        {where}
        GROUP BY strftime('%w', {col})
        ORDER BY num_dia
    """)


@flora.route("/api/flora/expedicao/pallet/cliente")
def flora_expedicao_pallet_cliente():
    return exp("""
        SELECT cliente, COALESCE(SUM(pallets), 0) AS total
        FROM b
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@flora.route("/api/flora/expedicao/pallet/destino")
def flora_expedicao_pallet_destino():
    return exp("""
        SELECT destino, COALESCE(SUM(pallets), 0) AS total
        FROM b
        {where}
        GROUP BY destino
        ORDER BY total DESC
    """)


@flora.route("/api/flora/expedicao/carga/por-cliente")
def flora_expedicao_carga_cliente():
    return exp("""
        SELECT cliente, COUNT(*) AS total
        FROM b
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@flora.route("/api/flora/expedicao/por-mes")
def flora_expedicao_por_mes():
    return exp("""
        SELECT strftime('%Y-%m', {col}) AS mes, COUNT(*) AS cargas,
               COALESCE(SUM(pallets), 0) AS total
        FROM b
        {where}
        GROUP BY strftime('%Y-%m', {col})
        ORDER BY mes
    """)


@flora.route("/api/flora/expedicao/permanencia/faixas")
def flora_expedicao_permanencia_faixas():
    # Quantas cargas ficaram até 7 dias, 8 a 15, 16 a 30 ou mais de 30 dias armazenadas
    return exp(f"""
        SELECT CASE ordem WHEN 1 THEN '0 a 7 dias'
                          WHEN 2 THEN '8 a 15 dias'
                          WHEN 3 THEN '16 a 30 dias'
                          ELSE 'Mais de 30 dias' END AS faixa,
               COUNT(*) AS total, SUM(pallets) AS pallets
        FROM (
            SELECT pallets,
                   CASE WHEN dias <= 7 THEN 1 WHEN dias <= 15 THEN 2 WHEN dias <= 30 THEN 3 ELSE 4 END AS ordem
            FROM (SELECT pallets, {PERMANENCIA} AS dias FROM b {{where}})
        )
        GROUP BY ordem
        ORDER BY ordem
    """)


@flora.route("/api/flora/expedicao/permanencia/cliente")
def flora_expedicao_permanencia_cliente():
    return exp(f"""
        SELECT cliente, COUNT(*) AS cargas,
               ROUND(AVG({PERMANENCIA}), 1) AS dias_medios
        FROM b
        {{where}}
        GROUP BY cliente
        ORDER BY dias_medios DESC
    """)


# ==========================
# API ESTOQUE (saldo = recebido - expedido; carga sem check_saida = em estoque)
# O dashboard chama estas rotas sem filtro de ano/mês (histórico completo).
# ==========================
@flora.route("/api/flora/estoque/por-cliente")
def flora_estoque_por_cliente():
    return rec("""
        SELECT cliente,
               SUM(pallets)                                          AS recebido,
               SUM(CASE WHEN saiu = 1 THEN pallets ELSE 0 END)       AS expedido,
               SUM(CASE WHEN saiu = 0 THEN pallets ELSE 0 END)       AS saldo,
               SUM(CASE WHEN saiu = 0 THEN 1 ELSE 0 END)             AS cargas_pendentes
        FROM b
        {where}
        GROUP BY cliente
        ORDER BY saldo DESC, recebido DESC
    """)


@flora.route("/api/flora/estoque/pendentes")
def flora_estoque_pendentes():
    # Cargas que ainda estão no armazém, da mais antiga para a mais nova
    return jsonify(consultar(BASE + """
        SELECT cliente, data_rec AS data, destino, pallets, nf AS nota_fiscal,
               CAST(julianday('now') - julianday(data_rec) AS INTEGER) AS dias
        FROM b
        WHERE saiu = 0
        ORDER BY data_rec, cliente
    """))


# ==========================
# EXPORTAR PARA EXCEL (período filtrado, mesma lista da aba Recebimento)
# ==========================
@flora.route("/api/flora/exportar")
def flora_exportar():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    where, params = filtro_dados("data_rec")
    linhas = consultar(BASE + f"""
        SELECT data_rec, cliente, nf, cte, pallets, destino, data_saida,
               CASE WHEN saiu = 1 THEN 'Expedido' ELSE 'Em estoque' END AS situacao,
               CASE WHEN saiu = 1 AND data_saida IS NOT NULL
                    THEN CAST(julianday(data_saida) - julianday(data_rec) AS INTEGER) END AS dias,
               observacao
        FROM b
        {where}
        ORDER BY data_rec, cliente
    """, params)

    cabecalho = ["Data recebimento", "Cliente", "Notas fiscais", "CT-e", "Pallets",
                 "Destino", "Data saída", "Situação", "Dias armazenado", "Observação"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Flora"
    ws.append(cabecalho)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="035A80")
        c.alignment = Alignment(horizontal="center")

    def dia(v):
        return datetime.strptime(v, "%Y-%m-%d").date() if v else None

    for l in linhas:
        ws.append([dia(l["data_rec"]), l["cliente"], l["nf"], l["cte"], l["pallets"], l["destino"],
                   dia(l["data_saida"]), l["situacao"], l["dias"], l["observacao"]])
    for linha in ws.iter_rows(min_row=2):
        for i in (0, 6):
            linha[i].number_format = "DD/MM/YYYY"
    for i, larg in enumerate([16, 26, 28, 10, 9, 20, 14, 13, 16, 30], start=1):
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    ano, mes = data()
    periodo = f"{ano or 'todos'}" + (f"-{mes}" if mes else "")
    return send_file(buf, as_attachment=True, download_name=f"flora_{periodo}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app = Flask(__name__)
    app.secret_key = "dev"  # só para testar este arquivo sozinho
    app.register_blueprint(flora)

    @app.route("/dev-login")
    def dev_login():  # só no teste: abre a sessão para as rotas /api/flora/... responderem
        session["usuario"], session["empresa"] = "teste", "flora"
        return "sessão aberta - agora acesse /api/flora/recebimento/resumo"


# ==========================================================
# API DE LANÇAMENTOS — FLORA
# Utiliza a tabela existente: in_out
# NÃO apaga nem recria os dados existentes.
# ==========================================================

from flask import current_app


# ----------------------------------------------------------
# CONEXÃO E CONSULTAS
# ----------------------------------------------------------

def conectar_lancamentos():
    conn = sqlite3.connect(CAMINHO_BANCO)
    conn.row_factory = sqlite3.Row
    return conn


def buscar_lancamento(conn, registro_id):
    return conn.execute("""
        SELECT
            rowid AS id,
            data_recebimento,
            cliente,
            no_cte,
            no_nota_fiscal,
            quantidade_paletes,
            rota_destino,
            COALESCE(check_saida, 0) AS check_saida,
            data_coleta_saida,
            observacao
        FROM in_out
        WHERE rowid = ?
    """, (registro_id,)).fetchone()


def lancamento_json(registro):
    if registro is None:
        return None

    return {
        "id": registro["id"],
        "data_recebimento": registro["data_recebimento"],
        "cliente": registro["cliente"] or "",
        "no_cte": registro["no_cte"],
        "no_nota_fiscal": registro["no_nota_fiscal"] or "",
        "quantidade_paletes": registro["quantidade_paletes"] or 0,
        "rota_destino": registro["rota_destino"] or "",
        "check_saida": registro["check_saida"],
        "data_coleta_saida": registro["data_coleta_saida"],
        "observacao": registro["observacao"] or ""
    }


def validar_data_iso(valor, campo):
    """Aceita datas no formato YYYY-MM-DD."""
    if not valor:
        return None, f"O campo '{campo}' é obrigatório."

    try:
        data_validada = datetime.strptime(str(valor), "%Y-%m-%d")
        if data_validada.strftime("%Y-%m-%d") != str(valor):
            raise ValueError
        return str(valor), None
    except (ValueError, TypeError):
        return None, f"A data '{campo}' deve estar no formato AAAA-MM-DD."


def validar_dados_recebimento(dados):
    """Valida os campos enviados pelo formulário."""

    data_recebimento, erro = validar_data_iso(
        dados.get("data_recebimento"),
        "data_recebimento"
    )
    if erro:
        return None, erro

    cliente = str(dados.get("cliente") or "").strip()
    nota_fiscal = str(dados.get("no_nota_fiscal") or "").strip()
    cte = str(dados.get("no_cte") or "").strip()
    destino = str(dados.get("rota_destino") or "").strip()
    observacao = str(dados.get("observacao") or "").strip()

    if not cliente:
        return None, "Informe o cliente."

    if len(cliente) > 200:
        return None, "O nome do cliente é muito longo."

    if not nota_fiscal:
        return None, "Informe pelo menos uma nota fiscal."

    if len(nota_fiscal) > 2000:
        return None, "O campo de notas fiscais é muito longo."

    if len(cte) > 100:
        return None, "O número do CT-e é muito longo."

    if len(destino) > 200:
        return None, "O destino é muito longo."

    if len(observacao) > 2000:
        return None, "A observação é muito longa."

    try:
        pallets_texto = str(
            dados.get("quantidade_paletes", 0)
        ).strip()

        if not pallets_texto:
            pallets_texto = "0"

        pallets = int(pallets_texto)

        if pallets < 0:
            return None, "A quantidade de pallets não pode ser negativa."

    except (ValueError, TypeError):
        return None, "A quantidade de pallets deve ser um número inteiro."

    return {
        "data_recebimento": data_recebimento,
        "cliente": cliente,
        "no_cte": cte or None,
        "no_nota_fiscal": nota_fiscal,
        "quantidade_paletes": pallets,
        "rota_destino": destino,
        "observacao": observacao
    }, None


def validar_id(registro_id):
    try:
        registro_id = int(registro_id)
        if registro_id <= 0:
            return None
        return registro_id
    except (ValueError, TypeError):
        return None


def limite_consulta():
    try:
        limite = int(request.args.get("limite", 200))
    except (ValueError, TypeError):
        limite = 200

    return max(1, min(limite, 1000))


def validar_data_saida(data_saida, data_recebimento):
    data_saida, erro = validar_data_iso(
        data_saida,
        "data_coleta_saida"
    )

    if erro:
        return None, erro

    if data_saida < str(data_recebimento)[:10]:
        return None, "A data de saída não pode ser anterior ao recebimento."

    return data_saida, None


# ----------------------------------------------------------
# RECEBIMENTO — LISTAR
# GET /api/flora/lancamentos/recebimento?limite=200
# ----------------------------------------------------------

@flora.route("/api/flora/lancamentos/recebimento", methods=["GET"])
def flora_lancamentos_listar_recebimento():
    conn = conectar_lancamentos()

    try:
        registros = conn.execute("""
            SELECT
                rowid AS id,
                data_recebimento,
                cliente,
                no_cte,
                no_nota_fiscal,
                quantidade_paletes,
                rota_destino,
                COALESCE(check_saida, 0) AS check_saida,
                data_coleta_saida,
                observacao
            FROM in_out
            WHERE data_recebimento IS NOT NULL
            ORDER BY date(data_recebimento) DESC, rowid DESC
            LIMIT ?
        """, (limite_consulta(),)).fetchall()

        return jsonify([lancamento_json(r) for r in registros])

    finally:
        conn.close()


# ----------------------------------------------------------
# RECEBIMENTO — CADASTRAR
# POST /api/flora/lancamentos/recebimento
# ----------------------------------------------------------

@flora.route("/api/flora/lancamentos/recebimento", methods=["POST"])
def flora_lancamentos_criar_recebimento():
    dados = request.get_json(silent=True) or {}

    campos, erro = validar_dados_recebimento(dados)

    if erro:
        return jsonify(erro=erro), 400

    conn = conectar_lancamentos()

    try:
        cursor = conn.execute("""
            INSERT INTO in_out (
                data_recebimento,
                cliente,
                no_cte,
                no_nota_fiscal,
                quantidade_paletes,
                rota_destino,
                check_saida,
                data_coleta_saida,
                observacao
            )
            VALUES (?, ?, ?, ?, ?, ?, 0, NULL, ?)
        """, (
            campos["data_recebimento"],
            campos["cliente"],
            campos["no_cte"],
            campos["no_nota_fiscal"],
            campos["quantidade_paletes"],
            campos["rota_destino"],
            campos["observacao"]
        ))

        novo_id = cursor.lastrowid
        conn.commit()

        registro = buscar_lancamento(conn, novo_id)

        return jsonify(
            mensagem="Recebimento cadastrado com sucesso.",
            lancamento=lancamento_json(registro)
        ), 201

    except sqlite3.Error:
        conn.rollback()
        current_app.logger.exception(
            "Erro ao cadastrar recebimento da Flora"
        )
        return jsonify(erro="Não foi possível cadastrar o recebimento."), 500

    finally:
        conn.close()


# ----------------------------------------------------------
# RECEBIMENTO — EDITAR
# PUT /api/flora/lancamentos/recebimento/<id>
# ----------------------------------------------------------

@flora.route(
    "/api/flora/lancamentos/recebimento/<int:registro_id>",
    methods=["PUT"]
)
def flora_lancamentos_editar_recebimento(registro_id):
    dados = request.get_json(silent=True) or {}

    campos, erro = validar_dados_recebimento(dados)

    if erro:
        return jsonify(erro=erro), 400

    conn = conectar_lancamentos()

    try:
        registro = buscar_lancamento(conn, registro_id)

        if registro is None:
            return jsonify(erro="Recebimento não encontrado."), 404

        if registro["check_saida"]:
            return jsonify(
                erro="Esta carga já foi expedida. Desfaça a expedição antes de editar o recebimento."
            ), 409

        conn.execute("""
            UPDATE in_out
            SET
                data_recebimento = ?,
                cliente = ?,
                no_cte = ?,
                no_nota_fiscal = ?,
                quantidade_paletes = ?,
                rota_destino = ?,
                observacao = ?
            WHERE rowid = ?
        """, (
            campos["data_recebimento"],
            campos["cliente"],
            campos["no_cte"],
            campos["no_nota_fiscal"],
            campos["quantidade_paletes"],
            campos["rota_destino"],
            campos["observacao"],
            registro_id
        ))

        conn.commit()

        registro = buscar_lancamento(conn, registro_id)

        return jsonify(
            mensagem="Recebimento atualizado com sucesso.",
            lancamento=lancamento_json(registro)
        )

    except sqlite3.Error:
        conn.rollback()
        current_app.logger.exception(
            "Erro ao editar recebimento da Flora"
        )
        return jsonify(erro="Não foi possível editar o recebimento."), 500

    finally:
        conn.close()


# ----------------------------------------------------------
# RECEBIMENTO — EXCLUIR
# DELETE /api/flora/lancamentos/recebimento/<id>
# ----------------------------------------------------------

@flora.route(
    "/api/flora/lancamentos/recebimento/<int:registro_id>",
    methods=["DELETE"]
)
def flora_lancamentos_excluir_recebimento(registro_id):
    conn = conectar_lancamentos()

    try:
        registro = buscar_lancamento(conn, registro_id)

        if registro is None:
            return jsonify(erro="Recebimento não encontrado."), 404

        if registro["check_saida"]:
            return jsonify(
                erro="Esta carga já foi expedida. Desfaça a expedição antes de excluir o recebimento."
            ), 409

        conn.execute("""
            DELETE FROM in_out
            WHERE rowid = ?
        """, (registro_id,))

        conn.commit()

        return jsonify(mensagem="Recebimento excluído com sucesso.")

    except sqlite3.Error:
        conn.rollback()
        current_app.logger.exception(
            "Erro ao excluir recebimento da Flora"
        )
        return jsonify(erro="Não foi possível excluir o recebimento."), 500

    finally:
        conn.close()


# ----------------------------------------------------------
# ESTOQUE — LISTAR CARGAS PENDENTES
# GET /api/flora/lancamentos/pendentes?limite=1000
# ----------------------------------------------------------

@flora.route("/api/flora/lancamentos/pendentes", methods=["GET"])
def flora_lancamentos_pendentes():
    conn = conectar_lancamentos()

    try:
        registros = conn.execute("""
            SELECT
                rowid AS id,
                data_recebimento,
                cliente,
                no_cte,
                no_nota_fiscal,
                quantidade_paletes,
                rota_destino,
                COALESCE(check_saida, 0) AS check_saida,
                data_coleta_saida,
                observacao
            FROM in_out
            WHERE data_recebimento IS NOT NULL
              AND COALESCE(check_saida, 0) = 0
            ORDER BY date(data_recebimento), rowid
            LIMIT ?
        """, (limite_consulta(),)).fetchall()

        return jsonify([lancamento_json(r) for r in registros])

    finally:
        conn.close()


# ----------------------------------------------------------
# EXPEDIÇÃO — LISTAR CARGAS EXPEDIDAS
# GET /api/flora/lancamentos/expedicao?limite=200
# ----------------------------------------------------------

@flora.route("/api/flora/lancamentos/expedicao", methods=["GET"])
def flora_lancamentos_listar_expedicao():
    conn = conectar_lancamentos()

    try:
        registros = conn.execute("""
            SELECT
                rowid AS id,
                data_recebimento,
                cliente,
                no_cte,
                no_nota_fiscal,
                quantidade_paletes,
                rota_destino,
                COALESCE(check_saida, 0) AS check_saida,
                data_coleta_saida,
                observacao
            FROM in_out
            WHERE data_recebimento IS NOT NULL
              AND COALESCE(check_saida, 0) = 1
            ORDER BY date(data_coleta_saida) DESC, rowid DESC
            LIMIT ?
        """, (limite_consulta(),)).fetchall()

        return jsonify([lancamento_json(r) for r in registros])

    finally:
        conn.close()


# ----------------------------------------------------------
# EXPEDIÇÃO — REGISTRAR SAÍDA
# POST /api/flora/lancamentos/expedicao/<id>
# Corpo JSON: {"data_coleta_saida": "2026-10-26"}
# ----------------------------------------------------------

@flora.route(
    "/api/flora/lancamentos/expedicao/<int:registro_id>",
    methods=["POST"]
)
def flora_lancamentos_registrar_saida(registro_id):
    dados = request.get_json(silent=True) or {}
    conn = conectar_lancamentos()

    try:
        registro = buscar_lancamento(conn, registro_id)

        if registro is None:
            return jsonify(erro="Carga não encontrada."), 404

        if registro["check_saida"]:
            return jsonify(erro="Esta carga já possui uma saída registrada."), 409

        data_saida, erro = validar_data_saida(
            dados.get("data_coleta_saida"),
            registro["data_recebimento"]
        )

        if erro:
            return jsonify(erro=erro), 400

        conn.execute("""
            UPDATE in_out
            SET check_saida = 1,
                data_coleta_saida = ?
            WHERE rowid = ?
        """, (data_saida, registro_id))

        conn.commit()

        registro = buscar_lancamento(conn, registro_id)

        return jsonify(
            mensagem="Saída registrada com sucesso.",
            lancamento=lancamento_json(registro)
        )

    except sqlite3.Error:
        conn.rollback()
        current_app.logger.exception(
            "Erro ao registrar saída da Flora"
        )
        return jsonify(erro="Não foi possível registrar a saída."), 500

    finally:
        conn.close()


# ----------------------------------------------------------
# EXPEDIÇÃO — CORRIGIR DATA DA SAÍDA
# PUT /api/flora/lancamentos/expedicao/<id>
# ----------------------------------------------------------

@flora.route(
    "/api/flora/lancamentos/expedicao/<int:registro_id>",
    methods=["PUT"]
)
def flora_lancamentos_editar_saida(registro_id):
    dados = request.get_json(silent=True) or {}
    conn = conectar_lancamentos()

    try:
        registro = buscar_lancamento(conn, registro_id)

        if registro is None:
            return jsonify(erro="Carga não encontrada."), 404

        if not registro["check_saida"]:
            return jsonify(
                erro="Esta carga ainda não possui uma saída registrada."
            ), 409

        data_saida, erro = validar_data_saida(
            dados.get("data_coleta_saida"),
            registro["data_recebimento"]
        )

        if erro:
            return jsonify(erro=erro), 400

        conn.execute("""
            UPDATE in_out
            SET data_coleta_saida = ?
            WHERE rowid = ?
        """, (data_saida, registro_id))

        conn.commit()

        registro = buscar_lancamento(conn, registro_id)

        return jsonify(
            mensagem="Data de saída atualizada com sucesso.",
            lancamento=lancamento_json(registro)
        )

    except sqlite3.Error:
        conn.rollback()
        current_app.logger.exception(
            "Erro ao editar saída da Flora"
        )
        return jsonify(erro="Não foi possível atualizar a saída."), 500

    finally:
        conn.close()


# ----------------------------------------------------------
# EXPEDIÇÃO — DESFAZER SAÍDA
# DELETE /api/flora/lancamentos/expedicao/<id>
# A carga volta a ficar pendente no estoque.
# ----------------------------------------------------------

@flora.route(
    "/api/flora/lancamentos/expedicao/<int:registro_id>",
    methods=["DELETE"]
)
def flora_lancamentos_desfazer_saida(registro_id):
    conn = conectar_lancamentos()

    try:
        registro = buscar_lancamento(conn, registro_id)

        if registro is None:
            return jsonify(erro="Carga não encontrada."), 404

        if not registro["check_saida"]:
            return jsonify(
                erro="Esta carga já está pendente no estoque."
            ), 409

        conn.execute("""
            UPDATE in_out
            SET check_saida = 0,
                data_coleta_saida = NULL
            WHERE rowid = ?
        """, (registro_id,))

        conn.commit()

        registro = buscar_lancamento(conn, registro_id)

        return jsonify(
            mensagem="Saída desfeita. A carga voltou para o estoque.",
            lancamento=lancamento_json(registro)
        )

    except sqlite3.Error:
        conn.rollback()
        current_app.logger.exception(
            "Erro ao desfazer saída da Flora"
        )
        return jsonify(erro="Não foi possível desfazer a saída."), 500

    finally:
        conn.close()


    app.run(debug=True)