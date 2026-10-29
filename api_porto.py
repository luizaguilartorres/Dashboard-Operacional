from flask import Flask, jsonify, request, Blueprint, session, current_app
import sqlite3
import re
import unicodedata
from datetime import datetime
from pathlib import Path

CAMINHO = Path(__file__).parent
(CAMINHO / "db").mkdir(exist_ok=True)
CAMINHO_BANCO = CAMINHO / "db" / "dados_base_porto.db"
CAMINHO_VEICULOS = CAMINHO / "db" / "dados_veiculos.db"   # vem de xlsx/veiculos.xlsx (placa + tipo)

porto = Blueprint("porto", __name__)


@porto.before_request
def proteger_api():
    if "usuario" not in session:
        return jsonify(erro="não autenticado"), 401
    if session.get("empresa") not in ("4log", "porto"):
        return jsonify(erro="sem permissão"), 403


def data():
    ano = request.args.get("ano")  # Ex.: /api/recebimento?ano=2026&mes=3 -> ano="2026"
    mes = request.args.get("mes")
    if mes:
        mes = mes.zfill(2)  # completa com zero à esquerda. Ex: 3 -> 03
    return ano, mes


def filtro_dados(coluna="data"):
    ano, mes = data()
    filtro, params = [], []
    if ano:
        filtro.append(f"strftime('%Y',{coluna}) = ?")
        params.append(ano)
    if mes:
        filtro.append(f"strftime('%m',{coluna}) = ?")
        params.append(mes)
    where = "WHERE " + " AND ".join(filtro) if filtro else ""
    return where, params


def consultar(sql, params):
    conn = sqlite3.connect(CAMINHO_BANCO)
    conn.row_factory = sqlite3.Row
    try:
        return [dict(l) for l in conn.execute(sql, params).fetchall()]
    finally:
        conn.close()


DIAS = """CASE CAST(strftime('%w', data) AS INTEGER)
            WHEN 0 THEN 'Domingo'
            WHEN 1 THEN 'Segunda'
            WHEN 2 THEN 'Terça'
            WHEN 3 THEN 'Quarta'
            WHEN 4 THEN 'Quinta'
            WHEN 5 THEN 'Sexta'
            WHEN 6 THEN 'Sábado'
            END"""


def rota(sql, coluna="data"):
    # No SQL use {where} para o filtro de ano/mês e {dias} para o nome do dia da semana.
    # Não coloque "f" antes das aspas: o .format() abaixo é quem preenche.
    where, params = filtro_dados(coluna)
    return jsonify(consultar(sql.format(where=where, dias=DIAS), params))


# ==========================
# API RECEBIMENTO-VEICULOS
# ==========================
@porto.route("/api/porto/recebimento/carro/por-hora")
def porto_recebimento_carro_por_hora():
    return rota("""
        SELECT strftime('%H', hora) AS hora, COUNT(DISTINCT num_cte) AS total
        FROM recebimento
        {where}
        GROUP BY strftime('%H', hora)
        ORDER BY hora
    """)


@porto.route("/api/porto/recebimento/carro/data")
def porto_recebimento_carro_por_data():
    return rota("""
        SELECT date(data) AS data, COUNT(DISTINCT num_cte) AS total
        FROM recebimento
        {where}
        GROUP BY date(data)
        ORDER BY data
    """)


@porto.route("/api/porto/recebimento/carro/dias-semana")
def porto_recebimento_carro_por_dias_semana():
    return rota("""
        SELECT strftime('%w', data) AS num_dia,
               {dias} AS dia_semana,
               COUNT(DISTINCT num_cte) AS total
        FROM recebimento
        {where}
        GROUP BY strftime('%w', data)
        ORDER BY num_dia
    """)


@porto.route("/api/porto/recebimento/carro/total")
def porto_recebimento_veiculos_recebidos():
    return rota("""
        SELECT COUNT(DISTINCT num_cte) AS total
        FROM recebimento
        {where}
    """)


@porto.route("/api/porto/recebimento/carro/por-tipo")
def porto_recebimento_tipos_de_veiculos():
    return rota("""
        SELECT tipo_do_veiculo AS veiculo, COUNT(DISTINCT num_cte) AS total
        FROM recebimento
        {where}
        GROUP BY tipo_do_veiculo
        ORDER BY total DESC
    """)


# ==========================
# API RECEBIMENTO-PALLETS
# ==========================
@porto.route("/api/porto/recebimento/pallet/total")
def porto_recebimento_total_pallet():
    return rota("""
        SELECT COALESCE(SUM(qtd_pallet), 0) AS total
        FROM recebimento
        {where}
    """)


@porto.route("/api/porto/recebimento/pallet/veiculo")
def porto_recebimento_pallet_veiculo():
    return rota("""
        SELECT tipo_do_veiculo AS veiculo, COALESCE(SUM(qtd_pallet), 0) AS total
        FROM recebimento
        {where}
        GROUP BY tipo_do_veiculo
        ORDER BY total DESC
    """)


@porto.route("/api/porto/recebimento/pallet/media-por-veiculo")
def porto_recebimento_pallet_media_veiculo():
    # Média de pallets por carro (viagem), não por linha da tabela.
    # "total" aqui é a média, para manter o mesmo nome de campo das outras rotas.
    return rota("""
        SELECT tipo_do_veiculo AS veiculo,
               ROUND(COALESCE(SUM(qtd_pallet), 0) * 1.0
                     / NULLIF(COUNT(DISTINCT num_cte), 0), 1) AS total
        FROM recebimento
        {where}
        GROUP BY tipo_do_veiculo
        ORDER BY total DESC
    """)


@porto.route("/api/porto/recebimento/pallet/cliente")
def porto_recebimento_pallet_cliente():
    return rota("""
        SELECT cliente, COALESCE(SUM(qtd_pallet), 0) AS total
        FROM recebimento
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@porto.route("/api/porto/recebimento/pallet/data")
def porto_recebimento_pallet_data():
    return rota("""
        SELECT date(data) AS data, COALESCE(SUM(qtd_pallet), 0) AS total
        FROM recebimento
        {where}
        GROUP BY date(data)
        ORDER BY data
    """)


@porto.route("/api/porto/recebimento/pallet/dia-semana")
def porto_recebimento_pallet_por_dia_semana():
    return rota("""
        SELECT strftime('%w', data) AS num_dia,
               {dias} AS dia_semana,
               COALESCE(SUM(qtd_pallet), 0) AS total
        FROM recebimento
        {where}
        GROUP BY strftime('%w', data)
        ORDER BY num_dia
    """)


# ==========================
# API RECEBIMENTO-PEDIDOS
# ==========================
@porto.route("/api/porto/recebimento/pedido/total")
def porto_recebimento_total_pedidos():
    return rota("""
        SELECT COUNT(DISTINCT num_pedido) AS total
        FROM recebimento
        {where}
    """)


@porto.route("/api/porto/recebimento/pedido/cliente")
def porto_recebimento_pedidos_por_cliente():
    return rota("""
        SELECT cliente, COUNT(DISTINCT num_pedido) AS total
        FROM recebimento
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@porto.route("/api/porto/recebimento/pedido/data")
def porto_recebimento_pedidos_por_data():
    return rota("""
        SELECT date(data) AS data, COUNT(DISTINCT num_pedido) AS total
        FROM recebimento
        {where}
        GROUP BY date(data)
        ORDER BY data
    """)


# =============================
# API RECEBIMENTO-NOTAS FISCAIS
# =============================
@porto.route("/api/porto/recebimento/nota-fiscal/total")
def porto_recebimento_total_notas():
    return rota("""
        SELECT COUNT(DISTINCT nota_fiscal) AS total
        FROM recebimento
        {where}
    """)


@porto.route("/api/porto/recebimento/nota-fiscal/data")
def porto_recebimento_notas_data():
    return rota("""
        SELECT date(data) AS data, COUNT(DISTINCT nota_fiscal) AS total
        FROM recebimento
        {where}
        GROUP BY date(data)
        ORDER BY data
    """)


# ==========================
# API EXPEDIÇÃO-VEICULOS
# ==========================
@porto.route("/api/porto/expedicao/carro/total")
def porto_expedicao_carro_total():
    return rota("""
        SELECT tipo_do_veiculo AS veiculo, COUNT(DISTINCT romaneio) AS total
        FROM expedicao
        {where}
        GROUP BY tipo_do_veiculo
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/carro/por-cliente")
def porto_expedicao_carro_por_cliente():
    return rota("""
        SELECT cliente, COUNT(DISTINCT romaneio) AS total
        FROM expedicao
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/carro/veiculo")
def porto_expedicao_carro_modelo():
    return rota("""
        SELECT tipo_do_veiculo AS veiculo, COUNT(DISTINCT romaneio) AS total
        FROM expedicao
        {where}
        GROUP BY tipo_do_veiculo
        ORDER BY total DESC
    """)


# ==========================
# API EXPEDIÇÃO-PALLETS
# ==========================
@porto.route("/api/porto/expedicao/pallet/total")
def porto_expedicao_pallet_total():
    return rota("""
        SELECT COALESCE(SUM(qtd_pallet), 0) AS total
        FROM expedicao
        {where}
    """)


@porto.route("/api/porto/expedicao/pallet/por-cliente")
def porto_expedicao_pallet_por_cliente():
    return rota("""
        SELECT cliente, COALESCE(SUM(qtd_pallet), 0) AS total
        FROM expedicao
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/pallet/por-destino")
def porto_expedicao_pallet_por_destino():
    return rota("""
        SELECT destino, COALESCE(SUM(qtd_pallet), 0) AS total
        FROM expedicao
        {where}
        GROUP BY destino
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/pallet/data")
def porto_expedicao_pallet_por_data():
    return rota("""
        SELECT date(data) AS data, COALESCE(SUM(qtd_pallet), 0) AS total
        FROM expedicao
        {where}
        GROUP BY date(data)
        ORDER BY data
    """)


# ==========================================================
# NOVAS APIs - RECEBIMENTO
# ==========================================================
@porto.route("/api/porto/recebimento/resumo")
def porto_recebimento_resumo():
    # Cards do topo: tudo em uma única chamada
    return rota("""
        SELECT COUNT(DISTINCT num_cte)                      AS carros,
               COALESCE(SUM(qtd_pallet), 0)             AS pallets,
               COUNT(DISTINCT num_pedido)               AS pedidos,
               COUNT(DISTINCT nota_fiscal)              AS notas,
               ROUND(COALESCE(SUM(qtd_pallet), 0) * 1.0
                     / NULLIF(COUNT(DISTINCT num_cte), 0), 1) AS media_pallet_carro
        FROM recebimento
        {where}
    """)


@porto.route("/api/porto/recebimento/carro/por-cliente")
def porto_recebimento_carro_por_cliente():
    return rota("""
        SELECT cliente, COUNT(DISTINCT num_cte) AS total
        FROM recebimento
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@porto.route("/api/porto/recebimento/carro/por-mes")
def porto_recebimento_carro_por_mes():
    return rota("""
        SELECT strftime('%Y-%m', data) AS mes, COUNT(DISTINCT num_cte) AS total
        FROM recebimento
        {where}
        GROUP BY strftime('%Y-%m', data)
        ORDER BY mes
    """)


@porto.route("/api/porto/recebimento/pallet/por-mes")
def porto_recebimento_pallet_por_mes():
    return rota("""
        SELECT strftime('%Y-%m', data) AS mes,
               COALESCE(SUM(qtd_pallet), 0) AS total
        FROM recebimento
        {where}
        GROUP BY strftime('%Y-%m', data)
        ORDER BY mes
    """)


@porto.route("/api/porto/recebimento/pallet/por-hora")
def porto_recebimento_pallet_por_hora():
    return rota("""
        SELECT strftime('%H', hora) AS hora,
               COALESCE(SUM(qtd_pallet), 0) AS total
        FROM recebimento
        {where}
        GROUP BY strftime('%H', hora)
        ORDER BY hora
    """)


# ==========================================================
# NOVAS APIs - EXPEDIÇÃO
# ==========================================================
@porto.route("/api/porto/expedicao/resumo")
def porto_expedicao_resumo():
    return rota("""
        SELECT COUNT(DISTINCT romaneio)                 AS carros,
               COALESCE(SUM(qtd_pallet), 0)             AS pallets,
               COALESCE(SUM(quantidade_de_caixas), 0)                 AS caixas,
               COUNT(DISTINCT num_pedido)               AS pedidos,
               COUNT(DISTINCT nota_fiscal)              AS notas,
               ROUND(COALESCE(SUM(quantidade_de_caixas), 0) * 1.0
                     / NULLIF(SUM(qtd_pallet), 0), 1)   AS caixas_por_pallet
        FROM expedicao
        {where}
    """)


# --- Veículos (romaneios) ---
@porto.route("/api/porto/expedicao/carro/quantidade")
def porto_expedicao_carro_quantidade():
    return rota("""
        SELECT COUNT(DISTINCT romaneio) AS total
        FROM expedicao
        {where}
    """)


@porto.route("/api/porto/expedicao/carro/data")
def porto_expedicao_carro_por_data():
    return rota("""
        SELECT date(data) AS data, COUNT(DISTINCT romaneio) AS total
        FROM expedicao
        {where}
        GROUP BY date(data)
        ORDER BY data
    """)


@porto.route("/api/porto/expedicao/carro/por-hora")
def porto_expedicao_carro_por_hora():
    return rota("""
        SELECT strftime('%H', hora) AS hora, COUNT(DISTINCT romaneio) AS total
        FROM expedicao
        {where}
        GROUP BY strftime('%H', hora)
        ORDER BY hora
    """)


@porto.route("/api/porto/expedicao/carro/dias-semana")
def porto_expedicao_carro_por_dias_semana():
    return rota("""
        SELECT strftime('%w', data) AS num_dia,
               {dias} AS dia_semana,
               COUNT(DISTINCT romaneio) AS total
        FROM expedicao
        {where}
        GROUP BY strftime('%w', data)
        ORDER BY num_dia
    """)


@porto.route("/api/porto/expedicao/carro/por-transportadora")
def porto_expedicao_carro_por_transportadora():
    return rota("""
        SELECT transportadora, COUNT(DISTINCT romaneio) AS total
        FROM expedicao
        {where}
        GROUP BY transportadora
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/carro/por-destino")
def porto_expedicao_carro_por_destino():
    return rota("""
        SELECT destino, COUNT(DISTINCT romaneio) AS total
        FROM expedicao
        {where}
        GROUP BY destino
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/carro/por-mes")
def porto_expedicao_carro_por_mes():
    return rota("""
        SELECT strftime('%Y-%m', data) AS mes, COUNT(DISTINCT romaneio) AS total
        FROM expedicao
        {where}
        GROUP BY strftime('%Y-%m', data)
        ORDER BY mes
    """)


# --- Pallets ---
@porto.route("/api/porto/expedicao/pallet/veiculo")
def porto_expedicao_pallet_veiculo():
    return rota("""
        SELECT tipo_do_veiculo AS veiculo, COALESCE(SUM(qtd_pallet), 0) AS total
        FROM expedicao
        {where}
        GROUP BY tipo_do_veiculo
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/pallet/media-por-veiculo")
def porto_expedicao_pallet_media_veiculo():
    # Média de pallets por carro (romaneio), não por linha da tabela
    return rota("""
        SELECT tipo_do_veiculo AS veiculo,
               ROUND(COALESCE(SUM(qtd_pallet), 0) * 1.0
                     / NULLIF(COUNT(DISTINCT romaneio), 0), 1) AS total
        FROM expedicao
        {where}
        GROUP BY tipo_do_veiculo
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/pallet/por-transportadora")
def porto_expedicao_pallet_por_transportadora():
    return rota("""
        SELECT transportadora, COALESCE(SUM(qtd_pallet), 0) AS total
        FROM expedicao
        {where}
        GROUP BY transportadora
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/pallet/dia-semana")
def porto_expedicao_pallet_por_dia_semana():
    return rota("""
        SELECT strftime('%w', data) AS num_dia,
               {dias} AS dia_semana,
               COALESCE(SUM(qtd_pallet), 0) AS total
        FROM expedicao
        {where}
        GROUP BY strftime('%w', data)
        ORDER BY num_dia
    """)


@porto.route("/api/porto/expedicao/pallet/por-mes")
def porto_expedicao_pallet_por_mes():
    return rota("""
        SELECT strftime('%Y-%m', data) AS mes,
               COALESCE(SUM(qtd_pallet), 0) AS total
        FROM expedicao
        {where}
        GROUP BY strftime('%Y-%m', data)
        ORDER BY mes
    """)


# --- Caixas ---
@porto.route("/api/porto/expedicao/caixa/total")
def porto_expedicao_caixa_total():
    return rota("""
        SELECT COALESCE(SUM(quantidade_de_caixas), 0) AS total
        FROM expedicao
        {where}
    """)


@porto.route("/api/porto/expedicao/caixa/por-cliente")
def porto_expedicao_caixa_por_cliente():
    return rota("""
        SELECT cliente, COALESCE(SUM(quantidade_de_caixas), 0) AS total
        FROM expedicao
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/caixa/por-destino")
def porto_expedicao_caixa_por_destino():
    return rota("""
        SELECT destino, COALESCE(SUM(quantidade_de_caixas), 0) AS total
        FROM expedicao
        {where}
        GROUP BY destino
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/caixa/data")
def porto_expedicao_caixa_por_data():
    return rota("""
        SELECT date(data) AS data, COALESCE(SUM(quantidade_de_caixas), 0) AS total
        FROM expedicao
        {where}
        GROUP BY date(data)
        ORDER BY data
    """)


@porto.route("/api/porto/expedicao/caixa/por-mes")
def porto_expedicao_caixa_por_mes():
    return rota("""
        SELECT strftime('%Y-%m', data) AS mes,
               COALESCE(SUM(quantidade_de_caixas), 0) AS total
        FROM expedicao
        {where}
        GROUP BY strftime('%Y-%m', data)
        ORDER BY mes
    """)


@porto.route("/api/porto/expedicao/caixa/por-pallet-cliente")
def porto_expedicao_caixa_por_pallet_cliente():
    # Caixas por pallet de cada cliente (indica o quanto o pallet vai "cheio")
    return rota("""
        SELECT cliente,
               ROUND(COALESCE(SUM(quantidade_de_caixas), 0) * 1.0
                     / NULLIF(SUM(qtd_pallet), 0), 1) AS total
        FROM expedicao
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


# --- Pedidos e notas ---
@porto.route("/api/porto/expedicao/pedido/total")
def porto_expedicao_pedido_total():
    return rota("""
        SELECT COUNT(DISTINCT num_pedido) AS total
        FROM expedicao
        {where}
    """)


@porto.route("/api/porto/expedicao/pedido/cliente")
def porto_expedicao_pedido_por_cliente():
    return rota("""
        SELECT cliente, COUNT(DISTINCT num_pedido) AS total
        FROM expedicao
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@porto.route("/api/porto/expedicao/nota-fiscal/total")
def porto_expedicao_nota_total():
    return rota("""
        SELECT COUNT(DISTINCT nota_fiscal) AS total
        FROM expedicao
        {where}
    """)


@porto.route("/api/porto/expedicao/nota-fiscal/data")
def porto_expedicao_nota_data():
    return rota("""
        SELECT date(data) AS data, COUNT(DISTINCT nota_fiscal) AS total
        FROM expedicao
        {where}
        GROUP BY date(data)
        ORDER BY data
    """)


# ==========================================================
# NOVAS APIs - DEVOLUÇÃO
# ==========================================================
@porto.route("/api/porto/devolucao/resumo")
def porto_devolucao_resumo():
    return rota("""
        SELECT COUNT(*)                                AS registros,
               COUNT(DISTINCT NULLIF(nota_fiscal, '-')) AS notas,
               COUNT(DISTINCT num_pedido)       AS pedidos,
               COALESCE(SUM(qtd_pallet), 0)     AS pallets,
               COALESCE(SUM(quantidade_de_caixas), 0)         AS caixas
        FROM devolucao
        {where}
    """)


@porto.route("/api/porto/devolucao/nota-fiscal/total")
def porto_devolucao_nota_total():
    return rota("""
        SELECT COUNT(DISTINCT NULLIF(nota_fiscal, '-')) AS total
        FROM devolucao
        {where}
    """)


@porto.route("/api/porto/devolucao/pallet/total")
def porto_devolucao_pallet_total():
    return rota("""
        SELECT COALESCE(SUM(qtd_pallet), 0) AS total
        FROM devolucao
        {where}
    """)


@porto.route("/api/porto/devolucao/caixa/total")
def porto_devolucao_caixa_total():
    return rota("""
        SELECT COALESCE(SUM(quantidade_de_caixas), 0) AS total
        FROM devolucao
        {where}
    """)


@porto.route("/api/porto/devolucao/por-motivo")
def porto_devolucao_por_motivo():
    return rota("""
        SELECT motivo,
               COUNT(*)  AS total,
               COALESCE(SUM(quantidade_de_caixas), 0)     AS caixas
        FROM devolucao
        {where}
        GROUP BY motivo
        ORDER BY total DESC
    """)


@porto.route("/api/porto/devolucao/por-situacao")
def porto_devolucao_por_situacao():
    return rota("""
        SELECT situacao,
               COUNT(*)  AS total,
               COALESCE(SUM(quantidade_de_caixas), 0)     AS caixas
        FROM devolucao
        {where}
        GROUP BY situacao
        ORDER BY total DESC
    """)


@porto.route("/api/porto/devolucao/por-cliente")
def porto_devolucao_por_cliente():
    return rota("""
        SELECT cliente,
               COUNT(*)  AS total,
               COALESCE(SUM(quantidade_de_caixas), 0)     AS caixas
        FROM devolucao
        {where}
        GROUP BY cliente
        ORDER BY total DESC
    """)


@porto.route("/api/porto/devolucao/por-transportadora")
def porto_devolucao_por_transportadora():
    return rota("""
        SELECT transportadora,
               COUNT(*)  AS total,
               COALESCE(SUM(quantidade_de_caixas), 0)     AS caixas
        FROM devolucao
        {where}
        GROUP BY transportadora
        ORDER BY total DESC
    """)


@porto.route("/api/porto/devolucao/data")
def porto_devolucao_por_data():
    return rota("""
        SELECT date(data) AS data,
               COUNT(*)  AS total,
               COALESCE(SUM(quantidade_de_caixas), 0)     AS caixas
        FROM devolucao
        {where}
        GROUP BY date(data)
        ORDER BY data
    """)


@porto.route("/api/porto/devolucao/por-mes")
def porto_devolucao_por_mes():
    return rota("""
        SELECT strftime('%Y-%m', data) AS mes,
               COUNT(*)  AS total,
               COALESCE(SUM(quantidade_de_caixas), 0)     AS caixas
        FROM devolucao
        {where}
        GROUP BY strftime('%Y-%m', data)
        ORDER BY mes
    """)


@porto.route("/api/porto/devolucao/lista")
def porto_devolucao_lista():
    # Linhas completas, para montar uma tabela no dashboard
    return rota("""
        SELECT date(data) AS data, hora, cliente, transportadora, nota_fiscal, nf_origem,
               num_pedido, motivo, situacao, qtd_pallet,
               quantidade_de_caixas AS caixas
        FROM devolucao
        {where}
        ORDER BY data DESC, hora DESC
    """)


# ==========================================================
# LANÇAMENTOS (inserir / listar / apagar / importar / exportar)
# ==========================================================
# Para ajustar os formulários, mexa aqui (campos aceitos) e em
# templates/lancamentos_porto.html (campos mostrados).
# Cada linha lançada guarda quem lançou (lancado_por) e quando (lancado_em).
LANC = {
    "recebimento": {
        "nome": "Recebimento",
        "cabecalho": ["data", "hora", "cliente", "transportadora", "num_cte", "placa", "tipo_do_veiculo"],
        "linha": ["nota_fiscal", "num_pedido", "qtd_pallet"],
        "obrigatorios": ["data", "cliente", "num_cte", "nota_fiscal"],
    },
    "expedicao": {
        "nome": "Expedição",
        "cabecalho": ["data", "hora", "cliente", "transportadora", "romaneio", "destino", "tipo_do_veiculo"],
        "linha": ["nota_fiscal", "num_pedido", "qtd_pallet", "quantidade_de_caixas"],
        "obrigatorios": ["data", "cliente", "romaneio", "nota_fiscal"],
    },
    "etiquetas": {
        "nome": "Etiquetas",
        # cabeçalho = o que se repete; linha = um pedido (ou um tamanho de etiqueta do pedido)
        "cabecalho": ["cliente", "solicitacao", "agenda"],
        "linha": ["pedido", "pallets", "caixas", "tamanho", "operador", "inicio", "conclusao", "etiquetadas"],
        "obrigatorios": ["pedido", "cliente", "solicitacao", "pallets", "caixas", "tamanho"],
        # nome do campo -> coluna real no banco (quando é diferente)
        "colunas": {"pedido": "numero_do_pedido", "solicitacao": "data_solicitacao", "agenda": "data_agenda",
                    "pallets": "quantidade_de_pallets", "caixas": "quantidade_de_caixas",
                    "tamanho": "tamanho_da_etiqueta", "inicio": "data_de_inicio",
                    "conclusao": "data_de_conclusao", "etiquetadas": "caixas_etiquetadas","cliente":"cliente"},
        # filtro de ano/mês da exportação
        "data_filtro": "COALESCE(data_de_conclusao, data_solicitacao)",
    },
    "devolucao": {
        "nome": "Devolução",
        "cabecalho": ["data", "hora", "cliente", "transportadora", "nota_fiscal", "nf_origem",
                      "num_pedido", "motivo", "situacao", "qtd_pallet", "quantidade_de_caixas"],
        "linha": [],
        "obrigatorios": ["data", "cliente"],
    },
}
TIPO_CAMPO = {
    "data": "date", "solicitacao": "date", "agenda": "date",
    "hora": "hora", "inicio": "datahora", "conclusao": "datahora",
    "qtd_pallet": "int", "quantidade_de_caixas": "int",   # num_cte e romaneio: texto livre
    "pallets": "int", "caixas": "int", "etiquetadas": "int",
    "cliente": "upper", "placa": "placa",
}
# sugestões (autocompletar): campo -> [(tabela, coluna), ...]
OPCOES = {
    "cliente": [("recebimento", "cliente"), ("expedicao", "cliente"), ("devolucao", "cliente"), ("etiquetas", "cliente")],
    "transportadora": [("recebimento", "transportadora"), ("expedicao", "transportadora"), ("devolucao", "transportadora")],
    "tipo_do_veiculo": [("recebimento", "tipo_do_veiculo"), ("expedicao", "tipo_do_veiculo")],
    "destino": [("expedicao", "destino")],
    "motivo": [("devolucao", "motivo")],
    "situacao": [("devolucao", "situacao")],
    "tamanho": [("etiquetas", "tamanho_da_etiqueta")],
    "operador": [("etiquetas", "operador")],
}
EXTRAS = {"lancado_por": "TEXT", "lancado_em": "TEXT"}
EXTRAS_TABELA = {"recebimento": {"placa": "TEXT"}}   # colunas extras só de algumas tabelas
PLACA_RE = re.compile(r"^[A-Z]{3}[0-9][A-Z0-9][0-9]{2}$")  # ABC1234 (antiga) ou ABC1D23 (Mercosul)
MAX_LINHAS_IMPORTACAO = 5000
MAX_BYTES_IMPORTACAO = 5 * 1024 * 1024


def _col(tab, campo):
    return LANC[tab].get("colunas", {}).get(campo, campo)


def _q(nome):
    return '"' + nome.replace('"', '""') + '"'


def _agora():
    try:
        from zoneinfo import ZoneInfo
        return datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _tabela_existe(conn, tabela):
    """Confere se a tabela principal já existe no banco importado."""
    return conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?",
        (tabela,),
    ).fetchone() is not None


def _garantir_colunas(conn, tabela):
    # Não tenta alterar uma tabela inexistente: isso gerava um erro SQLite
    # pouco claro durante a importação.
    if not _tabela_existe(conn, tabela):
        raise sqlite3.OperationalError(
            f"A tabela '{tabela}' não existe em {CAMINHO_BANCO}. "
            "Confira se a base foi criada/importada corretamente."
        )

    existentes = {
        r[1] for r in conn.execute(f"PRAGMA table_info({_q(tabela)})")
    }
    for c, tipo in {**EXTRAS, **EXTRAS_TABELA.get(tabela, {})}.items():
        if c not in existentes:
            conn.execute(f"ALTER TABLE {_q(tabela)} ADD COLUMN {_q(c)} {tipo}")


def _garantir_colunas_registro(conn, tabela, colunas):
    """Cria (como TEXT) as colunas que o lançamento grava e a tabela ainda não tem.
    Evita 'table X has no column named Y' quando o banco foi gerado antes de uma
    coluna nova existir. Os nomes vêm do próprio código (LANC), não do usuário."""
    existentes = {r[1] for r in conn.execute(f"PRAGMA table_info({_q(tabela)})")}
    for c in colunas:
        if c not in existentes:
            conn.execute(f"ALTER TABLE {_q(tabela)} ADD COLUMN {_q(c)} TEXT")


def _abrir(tabela):
    conn = sqlite3.connect(CAMINHO_BANCO)
    conn.row_factory = sqlite3.Row
    try:
        with conn:
            _garantir_colunas(conn, tabela)
            # Placas novas informadas nos lançamentos; a planilha de veículos fica intacta.
            conn.execute("""CREATE TABLE IF NOT EXISTS veiculos_extra (
                                placa TEXT PRIMARY KEY, tipo_do_veiculo TEXT,
                                atualizado_em TEXT, atualizado_por TEXT)""")
        return conn
    except Exception:
        conn.close()
        raise


def _usuario():
    return session.get("usuario") or ""


def _celula(v):
    # converte o que vem do Excel (datas, horas, números) em texto simples
    import math
    from datetime import date, time
    if v is None:
        return ""
    if isinstance(v, float) and math.isnan(v):
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return "" if s.lower() in ("nat", "nan", "none") else s


def _data(v):
    v = v.strip()
    for tam, fmt in ((10, "%Y-%m-%d"), (10, "%d/%m/%Y")):
        try:
            return datetime.strptime(v[:tam], fmt)
        except ValueError:
            pass
    raise ValueError("data inválida")


def remover_acentos(texto):
    # Mesma regra aplicada às planilhas em app.py, para que o que for lançado
    # pelo sistema fique igual ao que veio do Excel (ex.: Armazém -> Armazem).
    return "".join(c for c in unicodedata.normalize("NFD", texto)
                   if unicodedata.category(c) != "Mn")


def _valor(campo, bruto):
    v = ("" if bruto is None else str(bruto)).strip()
    tipo = TIPO_CAMPO.get(campo, "text")
    if tipo == "date":
        return _data(v).strftime("%Y-%m-%d 00:00:00") if v else None
    if tipo == "datahora":
        if not v:
            return None
        v = v.replace("T", " ")
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(v, fmt).strftime("%Y-%m-%d %H:%M:%S")
            except ValueError:
                pass
        raise ValueError("data/hora inválida")
    if tipo == "hora":
        if not v:
            return None
        v = v.split(".")[0]
        for fmt in ("%H:%M:%S", "%H:%M"):
            try:
                return datetime.strptime(v, fmt).strftime("%H:%M:%S")
            except ValueError:
                pass
        raise ValueError("hora inválida")
    if tipo == "int":
        if v == "":
            return None
        try:
            n = int(float(v.replace(",", ".")))
        except ValueError:
            raise ValueError(f"{campo} deve ser um número")
        if n < 0:
            raise ValueError(f"{campo} não pode ser negativo")
        return n
    if tipo == "placa":
        v = re.sub(r"[^A-Za-z0-9]", "", v).upper()
        if not v:
            return None
        if not PLACA_RE.match(v):
            raise ValueError(f"placa inválida ({v})")
        return v
    v = remover_acentos(v)
    if tipo == "upper":
        v = v.upper()
    return v[:120] if v else None


def _registro(tab, bruto):
    # bruto: {campo: texto}. Devolve {coluna_do_banco: valor}, já validado.
    cfg = LANC[tab]
    falta = [c for c in cfg["obrigatorios"] if str(bruto.get(c, "")).strip() == ""]
    if falta:
        raise ValueError("campos obrigatórios: " + ", ".join(falta))
    campos = cfg["cabecalho"] + cfg["linha"]
    rec = {campo: _valor(campo, bruto.get(campo)) for campo in campos}
    out = {_col(tab, c): v for c, v in rec.items()}
    if tab == "etiquetas":
        ini, fim = rec.get("inicio"), rec.get("conclusao")
        out["hora_inicio"] = (ini[11:] + ".000000") if ini else None
        out["hora_conclusao"] = (fim[11:] + ".000000") if fim else None
    return out


# Nomes aceitos para a coluna do tipo do veículo no cadastro de placas
TIPO_COLUNAS = ("tipo_do_veiculo", "tipo_de_veiculo", "tipo", "tipo_veiculo",
                "modelo", "modelo_do_veiculo", "veiculo", "categoria")
# Tabela de cadastro que não tem "veiculo"/"placa" no nome
TABELAS_CADASTRO = ("modelo_carro",)


def _norm_placa(p):
    return re.sub(r"[^A-Za-z0-9]", "", str(p or "")).upper()


def _tipo_cadastrado(conn, placa):
    """Procura o tipo do veículo: 1) placas informadas nos lançamentos (veiculos_extra),
    2) cadastro de veículos: tabela modelo_carro, ou qualquer tabela com 'veiculo'/'placa'
    no nome e uma coluna 'placa' (no banco principal ou em db/dados_veiculos.db)."""
    row = conn.execute("SELECT tipo_do_veiculo FROM veiculos_extra WHERE placa = ?", (placa,)).fetchone()
    if row and row[0]:
        return row[0]
    fontes = [(conn, False)]
    if CAMINHO_VEICULOS.exists():
        fontes.append((sqlite3.connect(CAMINHO_VEICULOS), True))
    try:
        for c, fechar in fontes:
            extras = ", ".join("'" + t + "'" for t in TABELAS_CADASTRO)
            tabelas = [r[0] for r in c.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND "
                f"(name LIKE '%veiculo%' OR name LIKE '%placa%' OR name IN ({extras}))")
                if r[0] != "veiculos_extra"]
            for t in tabelas:
                cols = {r[1] for r in c.execute(f"PRAGMA table_info({_q(t)})")}
                tcol = next((x for x in TIPO_COLUNAS if x in cols), None)
                if "placa" not in cols or tcol is None:
                    continue
                r = c.execute(
                    f"SELECT {_q(tcol)} FROM {_q(t)} WHERE {_q(tcol)} IS NOT NULL AND "
                    "REPLACE(REPLACE(UPPER(TRIM(placa)), '-', ''), ' ', '') = ? LIMIT 1", (placa,)).fetchone()
                if r and str(r[0]).strip():
                    return str(r[0]).strip()
    finally:
        for c, fechar in fontes:
            if fechar:
                c.close()
    return None


def _resolver_placas(conn, registros, quem, quando):
    """Com placa sem tipo: busca o tipo no cadastro (placa desconhecida -> pede o tipo).
    Com placa e tipo diferentes do cadastro: guarda em veiculos_extra para as próximas vezes."""
    novas = {}
    for rec in registros:
        placa, tipo = rec.get("placa"), rec.get("tipo_do_veiculo")
        if not placa:
            continue
        if not tipo:
            tipo = novas.get(placa) or _tipo_cadastrado(conn, placa)
            if not tipo:
                raise ValueError(f"placa {placa} não está no cadastro de veículos: informe o tipo do veículo")
            rec["tipo_do_veiculo"] = tipo
        elif novas.get(placa) != tipo and _tipo_cadastrado(conn, placa) != tipo:
            novas[placa] = tipo
    for placa, tipo in novas.items():
        conn.execute("""INSERT INTO veiculos_extra (placa, tipo_do_veiculo, atualizado_em, atualizado_por)
                        VALUES (?, ?, ?, ?)
                        ON CONFLICT(placa) DO UPDATE SET tipo_do_veiculo = excluded.tipo_do_veiculo,
                            atualizado_em = excluded.atualizado_em, atualizado_por = excluded.atualizado_por""",
                     (placa, tipo, quando, quem))


def _inserir(tab, registros):
    cols = list(registros[0].keys()) + list(EXTRAS.keys())
    sql = f"INSERT INTO {tab} ({', '.join(_q(c) for c in cols)}) VALUES ({', '.join('?' * len(cols))})"
    quem, quando = _usuario(), _agora()
    conn = _abrir(tab)
    try:
        with conn:
            _garantir_colunas_registro(conn, tab, cols)
            if tab == "recebimento":
                _resolver_placas(conn, registros, quem, quando)
            conn.executemany(sql, [tuple(r[c] for c in registros[0]) + (quem, quando) for r in registros])
    finally:
        conn.close()


@porto.route("/api/porto/lancamentos/opcoes")
def porto_lanc_opcoes():
    conn = sqlite3.connect(CAMINHO_BANCO)
    try:
        out = {}
        for campo, fontes in OPCOES.items():
            vals = set()
            for t, col in fontes:
                try:
                    vals.update(r[0] for r in conn.execute(
                        f"SELECT DISTINCT {_q(col)} FROM {t} WHERE {_q(col)} IS NOT NULL AND TRIM({_q(col)}) <> ''"))
                except sqlite3.OperationalError:
                    pass
            out[campo] = sorted(vals, key=lambda x: str(x).lower())
        return jsonify(out)
    finally:
        conn.close()


@porto.route("/api/porto/lancamentos/<tabela>", methods=["GET"])
def porto_lanc_listar(tabela):
    if tabela not in LANC:
        return jsonify(erro="tabela inválida"), 404
    cfg = LANC[tabela]
    campos = list(dict.fromkeys(cfg["cabecalho"] + cfg["linha"]))
    limite = min(max(request.args.get("limite", 40, type=int), 1), 500)
    sel = []
    for c in campos:
        col, tipo = _q(_col(tabela, c)), TIPO_CAMPO.get(c, "text")
        sel.append(f"date({col}) AS {c}" if tipo == "date" else
                   f"substr({col}, 1, 16) AS {c}" if tipo == "datahora" else f"{col} AS {c}")
    sql = (f"SELECT rowid AS id, {', '.join(sel)}, lancado_por, lancado_em "
           f"FROM {tabela} ORDER BY rowid DESC LIMIT ?")
    conn = _abrir(tabela)
    try:
        return jsonify([dict(r) for r in conn.execute(sql, [limite]).fetchall()])
    finally:
        conn.close()


@porto.route("/api/porto/lancamentos/<tabela>", methods=["POST"])
def porto_lanc_inserir(tabela):
    if tabela not in LANC:
        return jsonify(erro="tabela inválida"), 404
    cfg = LANC[tabela]
    corpo = request.get_json(silent=True)
    if not isinstance(corpo, dict):
        return jsonify(erro="envie JSON"), 400
    cab, linhas = corpo.get("cabecalho") or {}, corpo.get("linhas") or []
    if cfg["linha"]:
        if not isinstance(linhas, list) or not linhas or len(linhas) > 200:
            return jsonify(erro="informe de 1 a 200 linhas"), 400
    else:
        linhas = [{}]
    try:
        registros = [_registro(tabela, {**cab, **ln}) for ln in linhas]
    except (ValueError, TypeError) as e:
        return jsonify(erro=str(e)), 400
    try:
        _inserir(tabela, registros)
    except ValueError as e:
        return jsonify(erro=str(e)), 400
    return jsonify(ok=True, inseridos=len(registros)), 201


@porto.route("/api/porto/veiculos/<placa>")
def porto_veiculo_por_placa(placa):
    """Tipo do veículo cadastrado para a placa (usado pelo formulário de recebimento)."""
    p = re.sub(r"[^A-Za-z0-9]", "", placa).upper()
    if not PLACA_RE.match(p):
        return jsonify(placa=p, valida=False, tipo_do_veiculo=None)
    conn = _abrir("recebimento")
    try:
        tipo = _tipo_cadastrado(conn, p)
    finally:
        conn.close()
    return jsonify(placa=p, valida=True, tipo_do_veiculo=tipo)


@porto.route("/api/porto/lancamentos/<tabela>/<int:rowid>", methods=["PUT"])
def porto_lanc_editar(tabela, rowid):
    """Edita uma linha já lançada. Recebe {campo: valor} com os campos da linha
    (cabeçalho + linha). Quem lançou e quando lançou não mudam."""
    if tabela not in LANC:
        return jsonify(erro="tabela inválida"), 404
    corpo = request.get_json(silent=True)
    if not isinstance(corpo, dict):
        return jsonify(erro="envie JSON"), 400
    try:
        rec = _registro(tabela, corpo)
    except (ValueError, TypeError) as e:
        return jsonify(erro=str(e)), 400
    quem, quando = _usuario(), _agora()
    conn = _abrir(tabela)
    try:
        with conn:
            if tabela == "recebimento":
                _resolver_placas(conn, [rec], quem, quando)
            sets = ", ".join(f"{_q(c)} = ?" for c in rec)
            n = conn.execute(f"UPDATE {tabela} SET {sets} WHERE rowid = ?",
                             list(rec.values()) + [rowid]).rowcount
    except ValueError as e:
        return jsonify(erro=str(e)), 400
    finally:
        conn.close()
    return jsonify(ok=bool(n), editados=n), (200 if n else 404)


@porto.route("/api/porto/lancamentos/<tabela>/<int:rowid>", methods=["DELETE"])
def porto_lanc_apagar(tabela, rowid):
    if tabela not in LANC:
        return jsonify(erro="tabela inválida"), 404
    conn = sqlite3.connect(CAMINHO_BANCO)
    try:
        with conn:
            n = conn.execute(f"DELETE FROM {tabela} WHERE rowid = ?", (rowid,)).rowcount
    finally:
        conn.close()
    return jsonify(ok=bool(n), apagados=n), (200 if n else 404)


def _normalizar(nome):
    import re
    import unicodedata
    nome = unicodedata.normalize("NFKD", str(nome)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", nome.lower()).strip("_")


@porto.route("/api/porto/lancamentos/<tabela>/importar", methods=["POST"])
def porto_lanc_importar(tabela):
    """Importa um .xlsx (primeira aba, ou a aba com o nome da tabela). Tudo ou nada:
    se qualquer linha tiver erro, nada é gravado e os erros são listados."""
    if tabela not in LANC:
        return jsonify(erro="tabela inválida"), 404
    arq = request.files.get("arquivo")
    if arq is None or not (arq.filename or "").lower().endswith(".xlsx"):
        return jsonify(erro="envie um arquivo .xlsx"), 400
    dados = arq.read(MAX_BYTES_IMPORTACAO + 1)
    if len(dados) > MAX_BYTES_IMPORTACAO:
        return jsonify(erro="arquivo maior que 5 MB"), 400
    import io
    import pandas as pd
    try:
        abas = pd.read_excel(io.BytesIO(dados), sheet_name=None, dtype=object)
    except Exception:
        return jsonify(erro="não foi possível ler o arquivo (é um .xlsx válido?)"), 400
    if not abas:
        return jsonify(erro="arquivo sem abas"), 400
    escolhida = next((n for n in abas if _normalizar(n) == tabela), next(iter(abas)))
    df = abas[escolhida]
    if len(df) > MAX_LINHAS_IMPORTACAO:
        return jsonify(erro=f"máximo de {MAX_LINHAS_IMPORTACAO} linhas por importação"), 400

    cfg = LANC[tabela]
    campos = list(dict.fromkeys(cfg["cabecalho"] + cfg["linha"]))
    alias = {}
    for c in campos:
        alias[_normalizar(c)] = c
        alias[_normalizar(_col(tabela, c))] = c
    mapa = {col: alias[_normalizar(col)] for col in df.columns if _normalizar(col) in alias}
    faltando = [c for c in cfg["obrigatorios"] if c not in mapa.values()]
    if faltando:
        return jsonify(erro="colunas obrigatórias ausentes na planilha: " + ", ".join(
            _col(tabela, c) for c in faltando)), 400

    registros, erros = [], []
    for i, linha in enumerate(df.to_dict("records"), start=2):  # linha 1 = cabeçalho
        bruto = {mapa[k]: _celula(v) for k, v in linha.items() if k in mapa}
        if not any(bruto.values()):
            continue
        try:
            registros.append(_registro(tabela, bruto))
        except (ValueError, TypeError) as e:
            erros.append(f"linha {i}: {e}")
    if erros:
        return jsonify(erro=f"{len(erros)} linha(s) com problema; nada foi importado",
                       detalhes=erros[:15]), 400
    if not registros:
        return jsonify(erro="nenhuma linha para importar"), 400
    try:
        _inserir(tabela, registros)
    except ValueError as e:
        # Erros de validação (ex.: placa não cadastrada) não são erros HTTP 500.
        return jsonify(erro=f"{e}; nada foi importado"), 400
    except sqlite3.Error:
        # Guarda o traceback completo no error.log do servidor.
        current_app.logger.exception(
            "Erro SQLite durante importação da tabela %s (banco: %s)",
            tabela, CAMINHO_BANCO,
        )
        return jsonify(
            erro=(
                "Falha ao gravar no banco de dados. Verifique se a tabela "
                f"'{tabela}' existe e se possui as colunas esperadas. "
                "O detalhe técnico foi registrado no error.log."
            )
        ), 500
    except Exception:
        current_app.logger.exception(
            "Erro inesperado durante importação da tabela %s", tabela
        )
        return jsonify(
            erro="Erro inesperado durante a importação. Consulte o error.log."
        ), 500

    return jsonify(ok=True, inseridos=len(registros)), 201


# ==========================================================
# EXPORTAR PARA EXCEL (inclui quem lançou cada linha)
# ==========================================================
@porto.route("/api/porto/exportar")
def porto_exportar():
    """?tabelas=recebimento,expedicao,etiquetas,devolucao (padrão: todas) &ano= &mes="""
    import io
    import re
    from flask import send_file
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    pedidas = [t for t in (request.args.get("tabelas") or "").split(",") if t in LANC] or list(LANC)
    ano, mes = data()
    wb = Workbook()
    wb.remove(wb.active)
    eh_data = re.compile(r"^\d{4}-\d{2}-\d{2}( \d{2}:\d{2}:\d{2})?$")

    def seguro(v):  # evita que texto digitado vire fórmula no Excel
        if not isinstance(v, str) or not v:
            return v
        if v[0] in ("=", "+", "@") or (v[0] == "-" and len(v) > 1 and not v[1:].replace(".", "").isdigit()):
            return "'" + v
        return v

    for tab in pedidas:
        cfg = LANC[tab]
        coluna_data = cfg.get("data_filtro", "data")
        filtro, params = [], []
        if ano:
            filtro.append(f"strftime('%Y', {coluna_data}) = ?"); params.append(ano)
        if mes:
            filtro.append(f"strftime('%m', {coluna_data}) = ?"); params.append(mes)
        where = "WHERE " + " AND ".join(filtro) if filtro else ""
        conn = _abrir(tab)
        try:
            cur = conn.execute(f"SELECT * FROM {tab} {where} ORDER BY rowid", params)
            cabecalho = [d[0] for d in cur.description]
            linhas = cur.fetchall()
        finally:
            conn.close()
        ws = wb.create_sheet(cfg["nome"])
        ws.append(cabecalho)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for r in linhas:
            vals = []
            for nome, v in zip(cabecalho, tuple(r)):
                if isinstance(v, str) and nome.startswith("data") and eh_data.match(v):
                    v = datetime.strptime(v, "%Y-%m-%d %H:%M:%S" if len(v) > 10 else "%Y-%m-%d")
                else:
                    v = seguro(v)
                vals.append(v)
            ws.append(vals)
        for i, nome in enumerate(cabecalho, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(nome) + 2, 12), 34)
            if nome.startswith("data"):
                for cel in ws[get_column_letter(i)][1:]:
                    cel.number_format = "DD/MM/YYYY HH:MM" if "inicio" in nome or "conclus" in nome or nome == "lancado_em" else "DD/MM/YYYY"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    periodo = f"{ano or 'todos'}" + (f"-{mes}" if mes else "")
    nome = f"porto_{pedidas[0] if len(pedidas) == 1 else 'dados'}_{periodo}.xlsx"
    return send_file(buf, as_attachment=True, download_name=nome,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ==========================================================
# API ETIQUETAS (cobrança = etiquetas usadas)
# ==========================================================
# Regra de classificação de cada pedido de etiquetagem (ajuste aqui):
#  - ETIQ_POR_PALLET: etiquetas coladas em cada pallet (hoje, 4 por pallet).
#  - É pedido de PALLET quando o tamanho é um dos TAMANHOS_PALLET (ex.: 100X150)
#    ou quando o valor lançado em "caixas_etiquetadas" é exatamente
#    ETIQ_POR_PALLET x pallets (e diferente do total de caixas), que é o que a
#    planilha faz hoje com a Amazon.
#  - Os demais pedidos são de CAIXA: etiquetas usadas = caixas_etiquetadas.
#  - Só conta para cobrança o que já foi concluído (tem data de conclusão).
ETIQ_POR_PALLET = 4
TAMANHOS_PALLET = ("100X150",)

_TAM = ", ".join(f"'{t}'" for t in TAMANHOS_PALLET)
ETQ_BASE = f"""
WITH b AS (
    SELECT rowid AS id,
           numero_do_pedido          AS pedido,
           cliente,
           date(data_solicitacao)    AS solicitacao,
           date(data_agenda)         AS agenda,
           COALESCE(quantidade_de_pallets, 0) AS pallets,
           COALESCE(quantidade_de_caixas, 0)  AS caixas,
           tamanho_da_etiqueta       AS tamanho,
           operador,
           data_de_inicio            AS inicio,
           data_de_conclusao         AS conclusao,
           COALESCE(caixas_etiquetadas, 0)    AS etiq,
           CASE WHEN data_de_conclusao IS NOT NULL THEN 1 ELSE 0 END AS concluido,
           date(COALESCE(data_de_conclusao, data_solicitacao))        AS ref
    FROM etiquetas
), c AS (
    SELECT *,
           CASE WHEN tamanho IN ({_TAM})
                  OR (etiq > 0 AND etiq = pallets * {ETIQ_POR_PALLET} AND etiq <> caixas)
                THEN 1 ELSE 0 END AS e_pallet
    FROM b
), e AS (
    SELECT *,
           CASE WHEN concluido = 1 AND e_pallet = 1 THEN pallets * {ETIQ_POR_PALLET} ELSE 0 END AS etq_pallet,
           CASE WHEN concluido = 1 AND e_pallet = 0 THEN etiq ELSE 0 END                         AS etq_caixa
    FROM c
)
"""


def etq(sql, extra=None):
    # Use {where} para o filtro de ano/mês (aplicado sobre a data de conclusão;
    # pedidos pendentes usam a data de solicitação). "extra" é uma condição fixa
    # que se soma ao filtro.
    where, params = filtro_dados("ref")
    if extra:
        where = f"{where} AND ({extra})" if where else f"WHERE {extra}"
    return jsonify(consultar((ETQ_BASE + sql).format(where=where), params))


@porto.route("/api/porto/etiquetas/resumo")
def porto_etq_resumo():
    return etq("""
        SELECT COUNT(DISTINCT pedido)                       AS pedidos,
               COALESCE(SUM(concluido), 0)                  AS concluidos,
               COUNT(*) - COALESCE(SUM(concluido), 0)       AS pendentes,
               COALESCE(SUM(pallets), 0)                    AS pallets,
               COALESCE(SUM(caixas), 0)                     AS caixas,
               COALESCE(SUM(etq_caixa), 0)                  AS etiquetas_caixa,
               COALESCE(SUM(etq_pallet), 0)                 AS etiquetas_pallet,
               COALESCE(SUM(etq_caixa + etq_pallet), 0)     AS etiquetas_total
        FROM e
        {where}
    """)


@porto.route("/api/porto/etiquetas/por-dia")
def porto_etq_por_dia():
    return etq("""
        SELECT ref AS data, COUNT(*) AS pedidos,
               SUM(etq_caixa) AS etiquetas_caixa, SUM(etq_pallet) AS etiquetas_pallet,
               SUM(etq_caixa + etq_pallet) AS etiquetas_total
        FROM e
        {where}
        GROUP BY ref
        ORDER BY ref
    """)


@porto.route("/api/porto/etiquetas/por-mes")
def porto_etq_por_mes():
    return etq("""
        SELECT strftime('%Y-%m', ref) AS mes, COUNT(*) AS pedidos,
               SUM(etq_caixa) AS etiquetas_caixa, SUM(etq_pallet) AS etiquetas_pallet,
               SUM(etq_caixa + etq_pallet) AS etiquetas_total
        FROM e
        {where}
        GROUP BY strftime('%Y-%m', ref)
        ORDER BY mes
    """)


@porto.route("/api/porto/etiquetas/por-cliente")
def porto_etq_por_cliente():
    return etq("""
        SELECT cliente, COUNT(*) AS pedidos,
               SUM(pallets) AS pallets, SUM(caixas) AS caixas,
               SUM(etq_caixa) AS etiquetas_caixa, SUM(etq_pallet) AS etiquetas_pallet,
               SUM(etq_caixa + etq_pallet) AS etiquetas_total
        FROM e
        {where}
        GROUP BY cliente
        ORDER BY etiquetas_total DESC
    """)


@porto.route("/api/porto/etiquetas/por-tamanho")
def porto_etq_por_tamanho():
    return etq("""
        SELECT tamanho, SUM(etq_caixa) AS etiquetas_caixa, SUM(etq_pallet) AS etiquetas_pallet,
               SUM(etq_caixa + etq_pallet) AS etiquetas_total
        FROM e
        {where}
        GROUP BY tamanho
        ORDER BY etiquetas_total DESC
    """)


@porto.route("/api/porto/etiquetas/cobranca")
def porto_etq_cobranca():
    # Base da cobrança: etiquetas usadas por cliente e tamanho de etiqueta
    return etq("""
        SELECT cliente, tamanho,
               SUM(etq_caixa) AS etiquetas_caixa, SUM(etq_pallet) AS etiquetas_pallet,
               SUM(etq_caixa + etq_pallet) AS etiquetas_total
        FROM e
        {where}
        GROUP BY cliente, tamanho
        HAVING SUM(etq_caixa + etq_pallet) > 0
        ORDER BY cliente, tamanho
    """)


@porto.route("/api/porto/etiquetas/por-operador")
def porto_etq_por_operador():
    return etq("""
        SELECT COALESCE(operador, 'Sem operador') AS operador, COUNT(*) AS pedidos,
               SUM(etq_caixa) AS etiquetas_caixa, SUM(etq_pallet) AS etiquetas_pallet,
               SUM(etq_caixa + etq_pallet) AS etiquetas_total
        FROM e
        {where}
        GROUP BY operador
        ORDER BY etiquetas_total DESC
    """)


# Pedidos com horário real de início e conclusão (a planilha tem muitos com 00:00)
_DUR = """(julianday(conclusao) - julianday(inicio)) * 24.0"""
_HORARIO_OK = """concluido = 1 AND time(inicio) <> '00:00:00' AND julianday(conclusao) > julianday(inicio)
                 AND julianday(conclusao) - julianday(inicio) < 1"""


@porto.route("/api/porto/etiquetas/produtividade")
def porto_etq_produtividade():
    return etq(f"""
        SELECT operador, COUNT(*) AS pedidos,
               ROUND(SUM({_DUR}), 2) AS horas,
               SUM(etq_caixa + etq_pallet) AS etiquetas,
               ROUND(SUM(etq_caixa + etq_pallet) / NULLIF(SUM({_DUR}), 0), 1) AS etiquetas_por_hora
        FROM e
        {{where}}
        GROUP BY operador
        ORDER BY etiquetas_por_hora DESC
    """, extra=f"operador IS NOT NULL AND {_HORARIO_OK}")


@porto.route("/api/porto/etiquetas/tempo-medio")
def porto_etq_tempo_medio():
    return etq(f"""
        SELECT COUNT(*) AS pedidos_com_horario,
               ROUND(AVG({_DUR}) * 60, 1) AS minutos_medios,
               ROUND(AVG(julianday(conclusao) - julianday(solicitacao)), 1) AS dias_lead_time
        FROM e
        {{where}}
    """, extra=_HORARIO_OK)


@porto.route("/api/porto/etiquetas/no-prazo")
def porto_etq_no_prazo():
    return etq("""
        SELECT COALESCE(SUM(CASE WHEN date(conclusao) <= agenda THEN 1 ELSE 0 END), 0) AS no_prazo,
               COALESCE(SUM(CASE WHEN date(conclusao) >  agenda THEN 1 ELSE 0 END), 0) AS atrasados
        FROM e
        {where}
    """, extra="concluido = 1 AND agenda IS NOT NULL")


@porto.route("/api/porto/etiquetas/pendentes")
def porto_etq_pendentes():
    return etq("""
        SELECT pedido, cliente, agenda, tamanho, pallets, caixas,
               etiq AS caixas_etiquetadas
        FROM e
        {where}
        ORDER BY agenda, cliente
    """, extra="concluido = 0")


@porto.route("/api/porto/etiquetas/lista")
def porto_etq_lista():
    return etq("""
        SELECT pedido, cliente, solicitacao, agenda, date(conclusao) AS conclusao,
               operador, tamanho, pallets, caixas,
               CASE WHEN e_pallet = 1 THEN 'Pallet' ELSE 'Caixa' END AS tipo,
               etq_caixa AS etiquetas_caixa, etq_pallet AS etiquetas_pallet,
               CASE WHEN concluido = 1 THEN 'Concluído' ELSE 'Pendente' END AS situacao
        FROM e
        {where}
        ORDER BY ref DESC, cliente
    """)


if __name__ == "__main__":
    app = Flask(__name__)
    app.secret_key = "dev"  # só para testar este arquivo sozinho
    app.register_blueprint(porto)
    app.run(debug=True)